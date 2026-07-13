# -*- coding: utf-8 -*-
# Markdownから説明用PPTXスライドを自動生成するスクリプト
import os
import sys
import re
import subprocess
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

base_dir = r"C:\Users\yert1\Documents\agy\10_Medical\Patient-information"
pptx_dir = os.path.join(base_dir, "pptx")
OFFICECLI_PATH = r"C:\Users\yert1\Documents\agy\00_System\bin\officecli.exe"

if not os.path.exists(pptx_dir):
    os.makedirs(pptx_dir)


def run_officecli_checks(document_path):
    """生成したPowerPointの構造検証と問題検出を行う。検証失敗では生成を中断しない。"""
    if not os.path.isfile(OFFICECLI_PATH):
        print(f"WARNING: OfficeCLIが見つからないため検証を省略します: {OFFICECLI_PATH}")
        return

    checks = [
        ("OpenXMLスキーマ検証", [OFFICECLI_PATH, "validate", document_path, "--json"]),
        ("スライド問題の検出", [OFFICECLI_PATH, "view", document_path, "issues", "--json"]),
    ]
    for label, command in checks:
        result = subprocess.run(command, capture_output=True, text=True, encoding="utf-8")
        if result.returncode == 0:
            print(f"OfficeCLI {label}: 完了")
            if result.stdout.strip():
                print(result.stdout.strip())
        else:
            detail = result.stderr.strip() or result.stdout.strip() or "詳細情報なし"
            print(f"WARNING: OfficeCLI {label}に失敗しました: {detail}")

def parse_markdown_to_sections(md_text):
    lines = md_text.splitlines()
    
    # 1. フロントマターの除去
    start_idx = 0
    title = "説明資料"
    if len(lines) > 0 and lines[0].strip() == "---":
        for i in range(1, len(lines)):
            if lines[i].strip() == "---":
                start_idx = i + 1
                break
            m_title = re.match(r"^title:\s*(.*)$", lines[i])
            if m_title:
                title = m_title.group(1).strip().strip('"').strip("'")

    body_lines = lines[start_idx:]
    
    # H1 タイトルの抽出
    for line in body_lines:
        if line.startswith("# ") or line.startswith("# **"):
            title = line.replace("#", "").replace("**", "").strip()
            break

    # セクション分割
    sections = []
    current_section = {
        "title": "表紙",
        "slide_summary": "",
        "body_lines": []
    }
    
    is_first = True
    
    for line in body_lines:
        cleaned = line.strip()
        if not cleaned:
            continue
            
        if cleaned.startswith("# ") or cleaned.startswith("# **"):
            continue # H1はスキップ
            
        # H2見出し (セクションの境界)
        if cleaned.startswith("## ") and not cleaned.startswith("###"):
            h2_text = cleaned[3:].replace("**", "").strip()
            if not is_first:
                sections.append(current_section)
            is_first = False
            current_section = {
                "title": h2_text,
                "slide_summary": "",
                "body_lines": []
            }
            continue
            
        if re.match(r"^\{\{visual:\s*[a-zA-Z0-9_-]+\s*\}\}$", cleaned):
            # 旧原稿との互換性のため図版指定は読み飛ばす。
            continue
            
        # スライド要約タグ
        slide_summary_match = re.match(r"^\{\{slide_summary:\s*(.*?)\s*\}\}$", cleaned)
        if slide_summary_match:
            current_section["slide_summary"] = slide_summary_match.group(1).strip()
            continue
            
        # 本文ライン
        current_section["body_lines"].append(line)
        
    sections.append(current_section)
    return title, sections

def create_slide_deck(title, sections, output_path):
    prs = Presentation()
    # 16:9 ワイド画面の設定
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # テーマカラー
    teal_dark = RGBColor(0x0F, 0x76, 0x6E) # 主見出し
    teal_light = RGBColor(0xE6, 0xFF, 0xFA) # 背景グラデ用ベース
    ink_dark = RGBColor(0x1F, 0x29, 0x33) # 文字色
    
    # 1. 表紙スライドの作成 (白地またはライトブルー背景にレイアウト)
    slide_layout = prs.slide_layouts[6] # 空白レイアウト
    slide = prs.slides.add_slide(slide_layout)
    
    # タイトルの追加
    tx_box = slide.shapes.add_textbox(Inches(1.0), Inches(2.2), Inches(11.333), Inches(3.0))
    tf = tx_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "患者さん・ご家族への説明資料"
    p.font.name = "Yu Gothic"
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = teal_dark
    p.alignment = PP_ALIGN.LEFT
    
    p2 = tf.add_paragraph()
    p2.text = title
    p2.font.name = "Yu Gothic"
    p2.font.size = Pt(36)
    p2.font.bold = True
    p2.font.color.rgb = ink_dark
    p2.space_before = Pt(14)
    p2.alignment = PP_ALIGN.LEFT
    
    # 表紙のスピーカーノート
    notes_slide = slide.notes_slide
    notes_slide.notes_text_frame.text = f"【表紙】\n{title}\n当院での診断結果および今後のご紹介についてご説明します。"

    # 2. 各通常セクションのスライド作成
    for sec in sections:
        # 表紙セクションが「表紙」というタイトルで最初にある場合はスキップ
        if sec["title"] == "表紙" and not sec["slide_summary"] and not sec["body_lines"]:
            continue
            
        slide = prs.slides.add_slide(slide_layout)
        
        # スライドタイトル
        tx_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.7), Inches(1.0))
        tf = tx_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = sec["title"]
        p.font.name = "Yu Gothic"
        p.font.size = Pt(30)
        p.font.bold = True
        p.font.color.rgb = teal_dark
        
        # 本文・スライド要約を全幅で表示する。
        tx_content_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.7), Inches(4.8))
        tf_content = tx_content_box.text_frame
        tf_content.word_wrap = True
        
        # スライド要約をメインに表示
        if sec["slide_summary"]:
            p_sum = tf_content.paragraphs[0]
            p_sum.text = sec["slide_summary"]
            p_sum.font.name = "Yu Gothic"
            p_sum.font.size = Pt(22)
            p_sum.font.bold = True
            p_sum.font.color.rgb = ink_dark
            p_sum.line_spacing = 1.4
        else:
            # 要約がない場合は本文の箇条書き等の最初の数行を表示
            first_p = True
            for line in sec["body_lines"][:5]:
                clean_line = line.strip().replace("**", "").replace("* ", "").replace("- ", "")
                if not clean_line:
                    continue
                if first_p:
                    p_body = tf_content.paragraphs[0]
                    first_p = False
                else:
                    p_body = tf_content.add_paragraph()
                p_body.text = "• " + clean_line
                p_body.font.name = "Yu Gothic"
                p_body.font.size = Pt(18)
                p_body.font.color.rgb = ink_dark
                p_body.space_after = Pt(10)
                
        # スピーカーノートに元の詳細テキストを流し込む
        notes_slide = slide.notes_slide
        clean_body = []
        for line in sec["body_lines"]:
            # スライド要約などのメタ記述は除外
            if line.strip().startswith("{{") and line.strip().endswith("}}"):
                continue
            # 太字のマークなどを綺麗にする
            clean_body.append(line.replace("**", "").strip())
        
        notes_text = f"【{sec['title']}】\n" + "\n".join(clean_body)
        notes_slide.notes_text_frame.text = notes_text
        
    prs.save(output_path)
    print(f"Successfully generated PPTX at {output_path}")
    run_officecli_checks(output_path)

if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Usage: python convert_to_pptx.py <markdown_file>")
        sys.exit(1)
        
    md_file = sys.argv[1]
    if not os.path.exists(md_file):
        print(f"Error: {md_file} does not exist.")
        sys.exit(1)
        
    base_name = os.path.splitext(os.path.basename(md_file))[0]
    pptx_path = os.path.join(pptx_dir, f"{base_name}_説明スライド.pptx")
    
    with open(md_file, "r", encoding="utf-8") as f:
        md_text = f.read()
        
    title, sections = parse_markdown_to_sections(md_text)
    create_slide_deck(title, sections, pptx_path)
