#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
Patient Explainer Skill: Markdown to Word (.docx) Converter
患者説明書Markdownを印刷・配布に適したWord文書（.docx）に組版・変換するスクリプト。
"""

import sys
import subprocess
from pathlib import Path
import re

# 依存パッケージの自己修復・インポート
try:
    import docx
    from docx.shared import Pt, RGBColor, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import parse_xml, OxmlElement
    from docx.oxml.ns import nsdecls, qn
except ImportError:
    print("python-docx をインストール中...")
    try:
        subprocess.check_call([sys.executable, "-m", "pip", "install", "python-docx", "--break-system-packages"])
    except Exception:
        subprocess.check_call([sys.executable, "-m", "pip", "install", "python-docx"])
    import docx
    from docx.shared import Pt, RGBColor, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import parse_xml, OxmlElement
    from docx.oxml.ns import nsdecls, qn


# カラーパレット（医療・信頼のトーン: Evidence Navy）
COLOR_PRIMARY = RGBColor(0x1A, 0x36, 0x5D)     # #1A365D ディープネイビー
COLOR_SECONDARY = RGBColor(0x2B, 0x6C, 0xB0)   # #2B6CB0 スレートブルー
COLOR_TEXT = RGBColor(0x2D, 0x37, 0x48)        # #2D3748 チャコールグレー
COLOR_MUTED = RGBColor(0x71, 0x80, 0x96)       # #718096 グレー
HEX_CALLOUT_BG = "F7FAFC"
HEX_CALLOUT_BORDER = "CBD5E0"
HEX_ALERT_BG = "FFF5F5"
HEX_ALERT_BORDER = "FEB2B2"
HEX_HEADER_BG = "EDF2F7"


def set_eastasia_font(run, font_name="游ゴシック"):
    """日本語フォント（EastAsia）を明示的に設定する。"""
    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    rFonts.set(qn('w:eastAsia'), font_name)
    rFonts.set(qn('w:ascii'), font_name)
    rFonts.set(qn('w:hAnsi'), font_name)


def set_cell_background(cell, fill_hex):
    """表セルの背景色を設定する。"""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)


def set_cell_margins(cell, top=120, bottom=120, left=150, right=150):
    """表セルの内側余白を設定する。"""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)


def add_callout_box(doc, text, is_alert=False):
    """注意喚起・重要事項のコールアウト枠を作成する。"""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.autofit = False
    tbl.columns[0].width = Inches(6.5)

    cell = tbl.cell(0, 0)
    bg_hex = HEX_ALERT_BG if is_alert else HEX_CALLOUT_BG
    border_hex = HEX_ALERT_BORDER if is_alert else HEX_CALLOUT_BORDER
    set_cell_background(cell, bg_hex)
    set_cell_margins(cell, top=140, bottom=140, left=180, right=180)

    # 左ボーダー強調
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:left w:val="single" w:sz="24" w:space="0" w:color="{border_hex}"/>'
        f'<w:top w:val="single" w:sz="4" w:space="0" w:color="{border_hex}"/>'
        f'<w:right w:val="single" w:sz="4" w:space="0" w:color="{border_hex}"/>'
        f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="{border_hex}"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(tcBorders)

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15

    for line in text.split("\n"):
        line = line.strip()
        if not line:
            continue
        run = p.add_run(line + "\n")
        run.font.size = Pt(9.5)
        run.font.color.rgb = COLOR_TEXT
        set_eastasia_font(run)


def convert_markdown_to_docx(md_text: str, output_path: str):
    """Markdownテキストを解析し、整形されたWord文書を作成する。"""
    doc = docx.Document()

    # ページ余白の設定（印刷しやすい標準マージン）
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    lines = md_text.splitlines()
    in_table = False
    table_lines = []
    in_callout = False
    callout_lines = []
    callout_alert = False

    def flush_table():
        nonlocal in_table, table_lines
        if not table_lines:
            in_table = False
            return
        
        # テーブル行の解析
        rows_data = []
        for tl in table_lines:
            # 罫線行 (例: |---|---|) をスキップ
            if re.match(r'^\s*\|?\s*[-:]+[-| :]*\|?\s*$', tl):
                continue
            cols = [c.strip() for c in tl.strip().strip('|').split('|')]
            rows_data.append(cols)

        if rows_data:
            num_cols = max(len(r) for r in rows_data)
            tbl = doc.add_table(rows=len(rows_data), cols=num_cols)
            tbl.autofit = True

            for r_idx, row in enumerate(rows_data):
                for c_idx in range(num_cols):
                    val = row[c_idx] if c_idx < len(row) else ""
                    cell = tbl.cell(r_idx, c_idx)
                    cell.text = val
                    p = cell.paragraphs[0]
                    p.paragraph_format.space_before = Pt(3)
                    p.paragraph_format.space_after = Pt(3)
                    set_cell_margins(cell)

                    if r_idx == 0:
                        # ヘッダー行
                        set_cell_background(cell, HEX_HEADER_BG)
                        for run in p.runs:
                            run.font.bold = True
                            run.font.size = Pt(9.5)
                            run.font.color.rgb = COLOR_PRIMARY
                            set_eastasia_font(run)
                    else:
                        for run in p.runs:
                            run.font.size = Pt(9.5)
                            run.font.color.rgb = COLOR_TEXT
                            set_eastasia_font(run)

            doc.add_paragraph().paragraph_format.space_after = Pt(4)

        table_lines = []
        in_table = False

    def flush_callout():
        nonlocal in_callout, callout_lines, callout_alert
        if callout_lines:
            add_callout_box(doc, "\n".join(callout_lines), is_alert=callout_alert)
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
        callout_lines = []
        in_callout = False
        callout_alert = False

    for line in lines:
        stripped = line.strip()

        # テーブル判定
        if stripped.startswith('|') and stripped.endswith('|'):
            if in_callout:
                flush_callout()
            in_table = True
            table_lines.append(stripped)
            continue
        elif in_table:
            flush_table()

        # コールアウト判定 (> [!WARNING] など)
        if stripped.startswith('>'):
            content = re.sub(r'^>\s?', '', line).strip()
            if '[!WARNING]' in content or '[!CAUTION]' in content:
                callout_alert = True
                content = content.replace('[!WARNING]', '【警告】').replace('[!CAUTION]', '【注意】')
            elif '[!IMPORTANT]' in content:
                content = content.replace('[!IMPORTANT]', '【重要】')
            elif '[!NOTE]' in content:
                content = content.replace('[!NOTE]', '【ご案内】')
            in_callout = True
            callout_lines.append(content)
            continue
        elif in_callout:
            flush_callout()

        # 空行
        if not stripped:
            continue

        # 見出し1 (# タイトル)
        if stripped.startswith('# '):
            text = stripped[2:].strip()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(12)
            run = p.add_run(text)
            run.font.size = Pt(18)
            run.font.bold = True
            run.font.color.rgb = COLOR_PRIMARY
            set_eastasia_font(run)
            continue

        # 見出し2 (## セクション)
        if stripped.startswith('## '):
            text = stripped[3:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run(text)
            run.font.size = Pt(13)
            run.font.bold = True
            run.font.color.rgb = COLOR_PRIMARY
            set_eastasia_font(run)
            continue

        # 見出し3 (### サブセクション)
        if stripped.startswith('### '):
            text = stripped[4:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(text)
            run.font.size = Pt(11)
            run.font.bold = True
            run.font.color.rgb = COLOR_SECONDARY
            set_eastasia_font(run)
            continue

        # 箇条書き (- または *)
        if re.match(r'^[-*]\s+', stripped):
            text = re.sub(r'^[-*]\s+', '', stripped)
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            run = p.add_run(text)
            run.font.size = Pt(10)
            run.font.color.rgb = COLOR_TEXT
            set_eastasia_font(run)
            continue

        # 通常本文
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.2
        run = p.add_run(stripped)
        run.font.size = Pt(10)
        run.font.color.rgb = COLOR_TEXT
        set_eastasia_font(run)

    if in_table:
        flush_table()
    if in_callout:
        flush_callout()

    doc.save(output_path)
    print(f"Word文書を作成しました: {output_path}")


def main():
    if len(sys.argv) < 2:
        print("Usage: python convert_to_docx.py <input.md> [output.docx]")
        sys.exit(1)

    input_path = Path(sys.argv[1])
    if not input_path.exists():
        print(f"ファイルが見つかりません: {input_path}")
        sys.exit(1)

    output_path = sys.argv[2] if len(sys.argv) > 2 else str(input_path.with_suffix('.docx'))
    md_text = input_path.read_text(encoding='utf-8')
    convert_markdown_to_docx(md_text, output_path)


if __name__ == "__main__":
    main()
